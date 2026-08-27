#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
embed_pptx: 将图片整页贴合拼入目标 PPTX 文档的新页面中
自动检测空白版式、清除残留占位符，以母版宽高完全铺满画布。
"""

from __future__ import annotations

import argparse
from pathlib import Path
import sys
from typing import Any

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    print(
        "[FAIL] python-pptx 未安装，请先运行: pip install python-pptx",
        file=sys.stderr,
    )
    sys.exit(1)


def embed_image_to_pptx(pptx_path: Path, image_path: Path) -> Path:
    resolved_pptx = Path(pptx_path).resolve()
    resolved_img = Path(image_path).resolve()

    if not resolved_img.exists():
        print(f"[FAIL] 待插入图片不存在: {resolved_img}", file=sys.stderr)
        sys.exit(1)

    if resolved_pptx.exists():
        prs: Any = Presentation(str(resolved_pptx))
    else:
        prs = Presentation()
        # 默认 16:9 画布: 13.333 x 7.5 英寸
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

    # 动态探测空白版式
    blank_layout = None
    for layout in prs.slide_layouts:
        if len(layout.placeholders) == 0:
            blank_layout = layout
            break
    if blank_layout is None:
        blank_layout = prs.slide_layouts[0]

    # 添加新幻灯片
    slide = prs.slides.add_slide(blank_layout)

    # 清空所有占位符（避免残留 "Click to add title" 等文字框）
    for ph in list(slide.placeholders):
        sp = getattr(ph, "_element", None)
        if sp is not None:
            parent = sp.getparent()
            if parent is not None:
                parent.remove(sp)

    # 按照母版宽高 (0, 0, slide_width, slide_height) 整页贴合铺满
    _ = slide.shapes.add_picture(
        str(resolved_img),
        Pt(0),
        Pt(0),
        width=prs.slide_width,
        height=prs.slide_height,
    )

    resolved_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(resolved_pptx))
    slide_w_in = (
        prs.slide_width.inches
        if prs.slide_width and hasattr(prs.slide_width, "inches")
        else 13.33
    )
    slide_h_in = (
        prs.slide_height.inches
        if prs.slide_height and hasattr(prs.slide_height, "inches")
        else 7.5
    )
    print(
        f"[OK] 成功将图片拼入 {resolved_pptx} 作为第 {len(prs.slides)} 页 (画布尺寸: {slide_w_in:.2f}\" x {slide_h_in:.2f}\")"
    )
    return resolved_pptx


def main() -> None:
    parser = argparse.ArgumentParser(
        description="将生成的图片整页无缝贴合拼入 PPTX 文档"
    )
    _ = parser.add_argument(
        "--image", "-i", required=True, help="待插入的图片路径"
    )
    _ = parser.add_argument(
        "--pptx", "-p", required=True, help="目标 PPTX 文件路径"
    )

    args = parser.parse_args()
    _ = embed_image_to_pptx(Path(str(args.pptx)), Path(str(args.image)))


if __name__ == "__main__":
    main()
