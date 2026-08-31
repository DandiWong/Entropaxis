#!/usr/bin/env python3
"""
generate_deck.py: 极简高管汇报级 PPT 原生生成器
基于字阶、字重、留白与精准栅格对齐，零多余边框与底块。
"""

import sys
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Color Palette (Tokens)
C_DEEP_BLUE = RGBColor(27, 54, 93)     # #1B365D 主标题 / 核心重点
C_ACCENT_BLUE = RGBColor(10, 58, 118)  # #0A3A76 标签 / 分类
C_ORANGE = RGBColor(194, 65, 12)       # #C2410C 强调数据 / 痛点
C_TEXT_MAIN = RGBColor(31, 41, 55)     # #1F2937 正文深灰 / 重点
C_TEXT_BODY = RGBColor(55, 65, 81)     # #374151 正文描述
C_TEXT_MUTED = RGBColor(100, 116, 139) # #64748B 次级辅助说明
C_ARROW = RGBColor(14, 116, 144)       # #0E7490 流转箭头

FONT_FAMILY = "PingFang SC"

def add_header(slide, category: str, title: str, subtitle: str):
    """添加标准顶部三段式 Header（分类定位 + 结论型主标 + 解释型副标）"""
    # Category Tag
    tx_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.35))
    tf_cat = tx_cat.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category
    p_cat.font.name = FONT_FAMILY
    p_cat.font.size = Pt(13)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_ACCENT_BLUE

    # Main Title
    tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.6))
    tf_title = tx_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = FONT_FAMILY
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = C_DEEP_BLUE

    # Subtitle
    tx_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.45))
    tf_sub = tx_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(14.5)
    p_sub.font.color.rgb = C_TEXT_MUTED


def build_minimalist_deck(output_path: str):
    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    left_margin = Inches(0.8)

    # ----------------------------------------------------
    # SLIDE 1: 3-Metric 现状基线与痛点页
    # ----------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    add_header(slide1, 
               "01 · 现状基线与业务痛点",
               "当前一篇文章的发表基线：6个月、12万元、5—8轮修回",
               "完全依赖外部供应商代写，不仅周期长、成本高，更导致核心研究经验与修改记录持续流失")

    metrics = [
        ("12", " 万元", "单篇委外写作成本", "仅包含单次写作费用\n改投其他期刊需重复支付高额改写成本"),
        ("6", " 个月", "单篇总体发表周期", "供应商初稿约 1 个月\n内部拉锯修回约 3 个月 · 投稿准备约 2 个月"),
        ("5—8", " 轮", "多方邮件修回轮次", "医学、统计与作者跨多方反复沟通\n在数十个 Word 与邮件版本间拉锯损耗")
    ]
    col_w = Inches(3.6)
    col_gap = Inches(0.46)
    top_pos1 = Inches(2.25)

    for i, (num, unit, label, desc) in enumerate(metrics):
        c_left = left_margin + i * (col_w + col_gap)
        tb = slide1.shapes.add_textbox(c_left, top_pos1, col_w, Inches(2.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        r_num = p1.add_run()
        r_num.text = num
        r_num.font.name = FONT_FAMILY
        r_num.font.size = Pt(46)
        r_num.font.bold = True
        r_num.font.color.rgb = C_ORANGE if i == 2 else C_DEEP_BLUE

        r_unit = p1.add_run()
        r_unit.text = unit
        r_unit.font.name = FONT_FAMILY
        r_unit.font.size = Pt(20)
        r_unit.font.bold = True
        r_unit.font.color.rgb = C_ORANGE if i == 2 else C_DEEP_BLUE

        p2 = tf.add_paragraph()
        p2.space_before = Pt(8)
        p2.text = label
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_MAIN

        p3 = tf.add_paragraph()
        p3.space_before = Pt(8)
        p3.text = desc
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(13)
        p3.font.color.rgb = C_TEXT_MUTED

    # Bottom Takeaway (Slide 1)
    tb_bot1 = slide1.shapes.add_textbox(left_margin, Inches(5.15), Inches(11.7), Inches(1.8))
    tf_bot1 = tb_bot1.text_frame
    tf_bot1.word_wrap = True
    tf_bot1.margin_left = tf_bot1.margin_top = tf_bot1.margin_right = tf_bot1.margin_bottom = 0

    p_b1_title = tf_bot1.paragraphs[0]
    p_b1_title.text = "深层业务痛点分析"
    p_b1_title.font.name = FONT_FAMILY
    p_b1_title.font.size = Pt(14)
    p_b1_title.font.bold = True
    p_b1_title.font.color.rgb = C_ACCENT_BLUE

    p_b1_body = tf_bot1.add_paragraph()
    p_b1_body.space_before = Pt(6)
    p_b1_body.text = "关键问题不只是直接费用：大量临床试验成果转化效率受限；核心试验数据在外部流转存在合规风险；更重要的是，修改依据、审阅经验与质控规则随供应商流失，未能转化为集团内部的数字化沉淀。"
    p_b1_body.font.name = FONT_FAMILY
    p_b1_body.font.size = Pt(14.5)
    p_b1_body.font.color.rgb = C_TEXT_BODY

    # ----------------------------------------------------
    # SLIDE 2: 3-Pillar 业务定位与解法页
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2,
               "02 · 业务定位与解法",
               "从“外部委托代写”转向“内部受控协同生产”",
               "以临床试验材料为唯一事实源，建立受控的人机协同生产与质量把关闭环")

    pillars = [
        ("01", "事实源绝对锁定", "严格基于 Protocol / SAP / CSR / TLF 原文，严禁 AI 编造或向外发散试验数据；核心数据缺失显性标红报错，确保医学事实 100% 真实可信。"),
        ("02", "80 / 20 人机协同分工", "系统自动承担 80% 的重复性搬砖（文献初筛、段落起草、标准三线表生成、期刊排版）；专家专注 20% 核心医学判断与临床洞察。"),
        ("03", "专家终审与成果归属", "AI 仅生成候选草案并暴露风险，医学专家逐条采纳/驳回并拥有最终科学裁决权；论文成果 100% 署名与学术归属于业务团队。")
    ]
    top_pos2 = Inches(2.25)

    for i, (tag, p_title, body) in enumerate(pillars):
        c_left = left_margin + i * (col_w + col_gap)
        tb = slide2.shapes.add_textbox(c_left, top_pos2, col_w, Inches(2.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = tag
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = C_ACCENT_BLUE

        p2 = tf.add_paragraph()
        p2.space_before = Pt(6)
        p2.text = p_title
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(17)
        p2.font.bold = True
        p2.font.color.rgb = C_DEEP_BLUE

        p3 = tf.add_paragraph()
        p3.space_before = Pt(8)
        p3.text = body
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(13.5)
        p3.font.color.rgb = C_TEXT_BODY

    # Bottom (Slide 2)
    tb_bot2 = slide2.shapes.add_textbox(left_margin, Inches(5.15), Inches(11.7), Inches(1.8))
    tf_bot2 = tb_bot2.text_frame
    tf_bot2.word_wrap = True
    tf_bot2.margin_left = tf_bot2.margin_top = tf_bot2.margin_right = tf_bot2.margin_bottom = 0

    p_b2_title = tf_bot2.paragraphs[0]
    p_b2_title.text = "生产管理范式跃迁"
    p_b2_title.font.name = FONT_FAMILY
    p_b2_title.font.size = Pt(14)
    p_b2_title.font.bold = True
    p_b2_title.font.color.rgb = C_ACCENT_BLUE

    p_b2_body = tf_bot2.add_paragraph()
    p_b2_body.space_before = Pt(6)
    p_b2_body.text = "打破传统“黑盒代写 + 邮件拉锯”的离散模式：将材料理解、初稿起草、独立质控、专家复审与标准交付统一整合在集团内部受控环境中，全流程可审计、可回溯、可复制。"
    p_b2_body.font.name = FONT_FAMILY
    p_b2_body.font.size = Pt(14.5)
    p_b2_body.font.color.rgb = C_TEXT_BODY

    # ----------------------------------------------------
    # SLIDE 3: 6-Stage 业务流水线与流转箭头页
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3,
               "03 · 功能架构与业务流水线",
               "端到端业务流水线：前置质控与专家终审",
               "以临床试验材料为输入，经 6 阶段标准化受控流转，一键交付符合目标期刊规范的完整投稿包")

    stages = [
        ("① 材料解析", "确定性信息抽取", "• 解析 CSR / Protocol\n• 提取 PICO 与主要终点\n• 缺失数据显性标红报错\n• 锁定事实源绝对基线"),
        ("② 文献调研", "受控 PubMed 检索", "• 定向检索权威背景证据\n• 严禁发散本研究数据\n• 生成可核验 DOI 引用\n• 消除虚构引用风险"),
        ("③ 初稿撰写", "目标期刊 IMRaD", "• 严格基于材料分段起草\n• 精准控制段落篇幅预算\n• 自动匹配期刊结构规范\n• 生成标准医学三线表"),
        ("④ 独立质控", "隔离审查自动修订", "• 上下文隔离无偏见审查\n• 数值与统计一致性核验\n• 自动输出段内修改建议\n• 前置暴露潜在合规风险"),
        ("⑤ 专家复审", "Web 结构化工作台", "• 三栏结构化精准比对\n• 段内建议逐条采纳/驳回\n• 专家责任签署与留痕\n• 保留最终科学裁决权"),
        ("⑥ 投稿交付", "标准交付包输出", "• 排版规范 DOCX 稿件\n• Cover Letter 与声明\n• 质控清单与审计报告\n• 一键打包直接投稿")
    ]
    stage_w = Inches(1.68)
    arrow_w = Inches(0.32)
    top_pos3 = Inches(2.25)

    for i, (stg_title, stg_sub, stg_bullets) in enumerate(stages):
        curr_left = left_margin + i * (stage_w + arrow_w)
        tb = slide3.shapes.add_textbox(curr_left, top_pos3, stage_w, Inches(2.7))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = stg_title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(15.5)
        p1.font.bold = True
        p1.font.color.rgb = C_DEEP_BLUE

        p2 = tf.add_paragraph()
        p2.space_before = Pt(3)
        p2.text = stg_sub
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = C_ACCENT_BLUE

        p3 = tf.add_paragraph()
        p3.space_before = Pt(8)
        p3.text = stg_bullets
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(11.5)
        p3.font.color.rgb = C_TEXT_BODY

        if i < 5:
            arr_left = curr_left + stage_w + Inches(0.04)
            tb_arr = slide3.shapes.add_textbox(arr_left, top_pos3 + Inches(0.4), arrow_w, Inches(0.5))
            tf_arr = tb_arr.text_frame
            tf_arr.word_wrap = False
            tf_arr.margin_left = tf_arr.margin_top = tf_arr.margin_right = tf_arr.margin_bottom = 0
            p_arr = tf_arr.paragraphs[0]
            p_arr.alignment = PP_ALIGN.CENTER
            p_arr.text = "→"
            p_arr.font.name = FONT_FAMILY
            p_arr.font.size = Pt(20)
            p_arr.font.bold = True
            p_arr.font.color.rgb = C_ARROW

    # Bottom (Slide 3)
    tb_bot3 = slide3.shapes.add_textbox(left_margin, Inches(5.25), Inches(11.7), Inches(1.6))
    tf_bot3 = tb_bot3.text_frame
    tf_bot3.word_wrap = True
    tf_bot3.margin_left = tf_bot3.margin_top = tf_bot3.margin_right = tf_bot3.margin_bottom = 0

    p_b3_title = tf_bot3.paragraphs[0]
    p_b3_title.text = "核心风控与治理机制"
    p_b3_title.font.name = FONT_FAMILY
    p_b3_title.font.size = Pt(14)
    p_b3_title.font.bold = True
    p_b3_title.font.color.rgb = C_ACCENT_BLUE

    p_b3_body = tf_bot3.add_paragraph()
    p_b3_body.space_before = Pt(6)
    p_b3_body.text = "• 事实源唯一锁定：严禁向外发散或编造临床试验数据；缺失关键信息前置显性阻断。\n• 独立 Agent 视角隔离质控：上下文隔离审查，数值、统计、引用一致性硬规则拦截。\n• 人类专家拥有最终科学裁决权：严格遵循 ICMJE 伦理规范，每处修改可追溯，专家签字确认后方可导出。"
    p_b3_body.font.name = FONT_FAMILY
    p_b3_body.font.size = Pt(13.5)
    p_b3_body.font.color.rgb = C_TEXT_BODY

    # ----------------------------------------------------
    # SLIDE 4: 4-Value 价值对照与落地决策页
    # ----------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4,
               "04 · 试点价值与推进决策",
               "真实项目试点验证四维价值，据实确立推广决策",
               "以当前“6个月、12万元、5—8轮修回”为基线对照，通过 2~3 个真实临床试验建立决策证据")

    values = [
        ("01  速度", "1 个月 → 1 小时", "首版可审稿就绪时间", "首版可审稿草案从等待外部供应商 1 个月，大幅缩短至 1 小时内部生成就绪，彻底消除排期等待。"),
        ("02  质量", "减少 50%+ 返工", "消除数据与格式笔误", "独立质控前置拦截数值矛盾与格式疏漏，大幅减少内部医学与统计团队的多轮反复修回。"),
        ("03  治理", "100% 痕迹可溯", "责任链条清晰闭环", "每处段落修改的原始依据、采纳驳回记录与决策人全程归档留痕，严格符合医学发表伦理规范。"),
        ("04  复用", "沉淀集团资产", "专家经验数字化复用", "目标期刊排版模板、审阅经验与医学专业规则转化为组织级资产，实现跨管线、跨项目复用。")
    ]
    col_w4 = Inches(2.7)
    col_gap4 = Inches(0.3)
    top_pos4 = Inches(2.2)

    for i, (tag, stat, note, desc) in enumerate(values):
        c_left = left_margin + i * (col_w4 + col_gap4)
        tb = slide4.shapes.add_textbox(c_left, top_pos4, col_w4, Inches(2.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = tag
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = C_ACCENT_BLUE

        p2 = tf.add_paragraph()
        p2.space_before = Pt(4)
        p2.text = stat
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(20)
        p2.font.bold = True
        p2.font.color.rgb = C_ORANGE if i == 0 else C_DEEP_BLUE

        p3 = tf.add_paragraph()
        p3.space_before = Pt(4)
        p3.text = note
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(14)
        p3.font.bold = True
        p3.font.color.rgb = C_TEXT_MAIN

        p4 = tf.add_paragraph()
        p4.space_before = Pt(6)
        p4.text = desc
        p4.font.name = FONT_FAMILY
        p4.font.size = Pt(12)
        p4.font.color.rgb = C_TEXT_MUTED

    # Bottom (Slide 4)
    tb_bot4 = slide4.shapes.add_textbox(left_margin, Inches(4.75), Inches(11.7), Inches(2.3))
    tf_bot4 = tb_bot4.text_frame
    tf_bot4.word_wrap = True
    tf_bot4.margin_left = tf_bot4.margin_top = tf_bot4.margin_right = tf_bot4.margin_bottom = 0

    p_next_title = tf_bot4.paragraphs[0]
    p_next_title.text = "下一步试点推进计划与决策门禁"
    p_next_title.font.name = FONT_FAMILY
    p_next_title.font.size = Pt(14)
    p_next_title.font.bold = True
    p_next_title.font.color.rgb = C_ACCENT_BLUE

    p_steps = tf_bot4.add_paragraph()
    p_steps.space_before = Pt(6)
    p_steps.text = "1. 选定代表性项目：选取 2~3 个代表业务真实场景的临床试验项目（涵盖不同方案复杂度与期刊定位）启动受控试点。\n2. 跨专业联合评审：组织医学写作、临床统计与发表负责人按统一评估口径开展复审，客观检验速度与质量收益。\n3. 确立推广决策：用真实周期、质量、成本与风控对照数据，明确后续全集团推广范围与资源投入。"
    p_steps.font.name = FONT_FAMILY
    p_steps.font.size = Pt(13)
    p_steps.font.color.rgb = C_TEXT_BODY

    p_ask = tf_bot4.add_paragraph()
    p_ask.space_before = Pt(10)
    
    r_ask_tag = p_ask.add_run()
    r_ask_tag.text = "提请领导支持事项： "
    r_ask_tag.font.name = FONT_FAMILY
    r_ask_tag.font.size = Pt(14)
    r_ask_tag.font.bold = True
    r_ask_tag.font.color.rgb = C_DEEP_BLUE

    r_ask_content = p_ask.add_run()
    r_ask_content.text = "协调 2~3 个试点临床项目样本  ·  指定医学写作与统计骨干协同  ·  统筹内部合规算力与数据沙箱配置"
    r_ask_content.font.name = FONT_FAMILY
    r_ask_content.font.size = Pt(13.5)
    r_ask_content.font.bold = True
    r_ask_content.font.color.rgb = C_TEXT_MAIN

    p_out = Path(output_path).resolve()
    p_out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(p_out))
    print(f"[OK] 成功生成高管级极简汇报 PPT: {p_out}")
    return p_out


def main():
    parser = argparse.ArgumentParser(description="生成极简高管汇报级 PPT")
    parser.add_argument("--output", "-o", default="output/minimalist_deck.pptx", help="输出 PPTX 文件路径")
    args = parser.parse_args()
    build_minimalist_deck(args.output)


if __name__ == "__main__":
    main()
